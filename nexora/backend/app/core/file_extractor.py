"""
Nexora File Extractor & Multi-Format Parsing Service
Processes and extracts clean, structured content from multi-format files:
- Archives: .zip, .rar, .tar, .tar.gz, .tgz (with zip-bomb & path traversal protection)
- Documents: .pdf, .docx
- Data: .json, .csv, .yaml, .yml, .xml, .toml
- Code & Text: .py, .js, .ts, .html, .css, .cpp, .java, .rs, .go, .sql, .md, .txt (all programming languages)

Enforces strict 10 MB per-file size limit.
"""
from __future__ import annotations

import io
import os
import re
import shutil
import tarfile
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

# Strict Limit: 10 Megabytes
MAX_FILE_SIZE_BYTES = 10 * 1024 * 1024  # 10 MB
MAX_DECOMPRESSED_ARCHIVE_BYTES = 50 * 1024 * 1024  # 50 MB
MAX_ARCHIVE_ENTRIES = 250
MAX_PREVIEW_CHARS = 12000

# Extensions categorized
CODE_EXTENSIONS = {
    ".py", ".pyw", ".js", ".mjs", ".cjs", ".ts", ".tsx", ".jsx",
    ".html", ".htm", ".css", ".scss", ".sass", ".less",
    ".c", ".h", ".cpp", ".hpp", ".cc", ".hh", ".cxx", ".cs",
    ".go", ".rs", ".java", ".kt", ".kts", ".scala", ".rb",
    ".php", ".sh", ".bash", ".zsh", ".ps1", ".bat", ".cmd",
    ".sql", ".r", ".swift", ".dart", ".lua", ".zig", ".v",
    ".proto", ".graphql", ".gql", ".sol"
}

DATA_EXTENSIONS = {
    ".json", ".jsonl", ".csv", ".tsv", ".yaml", ".yml",
    ".xml", ".toml", ".ini", ".cfg", ".conf", ".env", ".properties"
}

TEXT_EXTENSIONS = {
    ".txt", ".md", ".markdown", ".rst", ".log", ".tex", ".asc"
}

ARCHIVE_EXTENSIONS = {
    ".zip", ".rar", ".tar", ".gz", ".tgz"
}

DOCUMENT_EXTENSIONS = {
    ".pdf", ".docx", ".doc"
}


@dataclass
class ExtractedFileResult:
    filename: str
    original_size: int
    category: str  # 'code', 'data', 'text', 'document', 'archive', 'binary'
    mime_hint: str
    text_content: str = ""
    summary: str = ""
    extracted_files: List[str] = field(default_factory=list)
    sandbox_path: Optional[str] = None
    page_count: Optional[int] = None
    error: Optional[str] = None

    def to_dict(self) -> Dict[str, Any]:
        return {
            "filename": self.filename,
            "original_size": self.original_size,
            "category": self.category,
            "mime_hint": self.mime_hint,
            "text_content": self.text_content,
            "summary": self.summary,
            "extracted_files": self.extracted_files,
            "sandbox_path": self.sandbox_path,
            "page_count": self.page_count,
            "error": self.error,
        }


def get_file_category(filename: str) -> Tuple[str, str]:
    """Returns (category, extension) based on filename."""
    lower = filename.lower()
    if lower.endswith(".tar.gz"):
        return "archive", ".tar.gz"
    
    ext = Path(lower).suffix
    if ext in CODE_EXTENSIONS:
        return "code", ext
    if ext in DATA_EXTENSIONS:
        return "data", ext
    if ext in TEXT_EXTENSIONS:
        return "text", ext
    if ext in ARCHIVE_EXTENSIONS:
        return "archive", ext
    if ext in DOCUMENT_EXTENSIONS:
        return "document", ext
    return "text", ext or ".txt"


def decode_text_bytes(data: bytes) -> str:
    """Attempts multi-encoding decode for text/code data."""
    for enc in ("utf-8", "utf-8-sig", "latin-1", "cp1252", "iso-8859-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def extract_pdf_text(data: bytes, filename: str) -> ExtractedFileResult:
    """Extracts text content and pages from a PDF file."""
    try:
        import pypdf
        reader = pypdf.PdfReader(io.BytesIO(data))
        page_count = len(reader.pages)
        pages_text: List[str] = []

        for idx, page in enumerate(reader.pages):
            text = (page.extract_text() or "").strip()
            if text:
                pages_text.append(f"--- Page {idx + 1} ---\n{text}")

        full_text = "\n\n".join(pages_text) if pages_text else "[PDF contains no selectable text or scanned images]"
        summary = f"PDF document with {page_count} page(s), extracted {len(full_text)} characters."

        return ExtractedFileResult(
            filename=filename,
            original_size=len(data),
            category="document",
            mime_hint="application/pdf",
            text_content=full_text,
            summary=summary,
            page_count=page_count,
        )
    except Exception as exc:
        return ExtractedFileResult(
            filename=filename,
            original_size=len(data),
            category="document",
            mime_hint="application/pdf",
            error=f"PDF parsing error: {exc}",
            summary="Failed to parse PDF document.",
        )


def extract_docx_text(data: bytes, filename: str) -> ExtractedFileResult:
    """Extracts text, headings, and tables from a Word DOCX file."""
    try:
        import docx
        doc = docx.Document(io.BytesIO(data))
        lines: List[str] = []

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            if para.style and para.style.name.startswith("Heading"):
                lines.append(f"### {text}")
            else:
                lines.append(text)

        # Extract tables
        for t_idx, table in enumerate(doc.tables):
            lines.append(f"\n[Table {t_idx + 1}]")
            for row in table.rows:
                cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                lines.append("| " + " | ".join(cells) + " |")

        full_text = "\n\n".join(lines) if lines else "[Word document is empty]"
        summary = f"DOCX document with {len(doc.paragraphs)} paragraphs and {len(doc.tables)} tables."

        return ExtractedFileResult(
            filename=filename,
            original_size=len(data),
            category="document",
            mime_hint="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            text_content=full_text,
            summary=summary,
        )
    except Exception as exc:
        return ExtractedFileResult(
            filename=filename,
            original_size=len(data),
            category="document",
            mime_hint="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            error=f"DOCX parsing error: {exc}",
            summary="Failed to parse DOCX document.",
        )


def _is_safe_path(base_dir: Path, target_path: Path) -> bool:
    """Verifies target path resides strictly inside base_dir (prevents zip-slip)."""
    try:
        resolved_base = base_dir.resolve()
        resolved_target = target_path.resolve()
        return resolved_base in resolved_target.parents or resolved_base == resolved_target
    except Exception:
        return False


def extract_archive(
    data: bytes,
    filename: str,
    output_dir: Path,
) -> ExtractedFileResult:
    """Safely extracts ZIP or TAR archive with zip-bomb and path-traversal safeguards."""
    output_dir.mkdir(parents=True, exist_ok=True)
    extracted_names: List[str] = []
    total_uncompressed_bytes = 0
    lower = filename.lower()

    try:
        if lower.endswith(".zip"):
            with zipfile.ZipFile(io.BytesIO(data), "r") as zf:
                info_list = zf.infolist()
                if len(info_list) > MAX_ARCHIVE_ENTRIES:
                    raise ValueError(f"Archive contains {len(info_list)} items, exceeding limit of {MAX_ARCHIVE_ENTRIES}.")

                for member in info_list:
                    # Zip-bomb protection check
                    total_uncompressed_bytes += member.file_size
                    if total_uncompressed_bytes > MAX_DECOMPRESSED_ARCHIVE_BYTES:
                        raise ValueError(f"Archive uncompressed size exceeds safe limit ({MAX_DECOMPRESSED_ARCHIVE_BYTES // (1024*1024)} MB).")

                    # Path traversal sanitization
                    member_path = Path(member.filename)
                    if member_path.is_absolute() or ".." in member_path.parts:
                        continue  # Skip unsafe entry

                    dest = output_dir / member_path
                    if not _is_safe_path(output_dir, dest):
                        continue

                    if member.is_dir():
                        dest.mkdir(parents=True, exist_ok=True)
                    else:
                        dest.parent.mkdir(parents=True, exist_ok=True)
                        with zf.open(member) as src, open(dest, "wb") as dst:
                            shutil.copyfileobj(src, dst)
                        extracted_names.append(member.filename)

        elif lower.endswith((".tar", ".tar.gz", ".tgz", ".gz")):
            with tarfile.open(fileobj=io.BytesIO(data), mode="r:*") as tf:
                members = tf.getmembers()
                if len(members) > MAX_ARCHIVE_ENTRIES:
                    raise ValueError(f"Archive contains {len(members)} items, exceeding limit of {MAX_ARCHIVE_ENTRIES}.")

                for member in members:
                    total_uncompressed_bytes += member.size
                    if total_uncompressed_bytes > MAX_DECOMPRESSED_ARCHIVE_BYTES:
                        raise ValueError(f"Archive uncompressed size exceeds safe limit.")

                    dest = output_dir / member.name
                    if not _is_safe_path(output_dir, dest):
                        continue

                    if member.isdir():
                        dest.mkdir(parents=True, exist_ok=True)
                    else:
                        dest.parent.mkdir(parents=True, exist_ok=True)
                        f = tf.extractfile(member)
                        if f:
                            with open(dest, "wb") as dst:
                                shutil.copyfileobj(f, dst)
                            extracted_names.append(member.name)
        elif lower.endswith(".rar"):
            # RAR support via rarfile if available, otherwise clear guidance
            try:
                import rarfile
                with rarfile.RarFile(io.BytesIO(data), "r") as rf:
                    for member in rf.infolist():
                        total_uncompressed_bytes += member.file_size
                        if total_uncompressed_bytes > MAX_DECOMPRESSED_ARCHIVE_BYTES:
                            raise ValueError("RAR uncompressed size exceeds safe limit.")
                        dest = output_dir / member.filename
                        if not _is_safe_path(output_dir, dest):
                            continue
                        if member.isdir():
                            dest.mkdir(parents=True, exist_ok=True)
                        else:
                            dest.parent.mkdir(parents=True, exist_ok=True)
                            with rf.open(member) as src, open(dest, "wb") as dst:
                                shutil.copyfileobj(src, dst)
                            extracted_names.append(member.filename)
            except Exception as rar_err:
                # If unrar tool is missing on system, record informative note
                return ExtractedFileResult(
                    filename=filename,
                    original_size=len(data),
                    category="archive",
                    mime_hint="application/x-rar-compressed",
                    error=f"RAR archive received: {rar_err}. Please upload as .zip or install unrar on host.",
                    summary="RAR archive requires host unrar binary; recommend .zip format.",
                )

        # Build tree summary
        tree_lines = [f"Archive `{filename}` unpacked into workspace ({len(extracted_names)} files):"]
        for name in extracted_names[:30]:
            tree_lines.append(f"  - {name}")
        if len(extracted_names) > 30:
            tree_lines.append(f"  ... and {len(extracted_names) - 30} more files.")

        summary_text = "\n".join(tree_lines)
        return ExtractedFileResult(
            filename=filename,
            original_size=len(data),
            category="archive",
            mime_hint="application/zip",
            text_content=summary_text,
            summary=f"Extracted {len(extracted_names)} files safely ({total_uncompressed_bytes} bytes unpacked).",
            extracted_files=extracted_names,
            sandbox_path=str(output_dir),
        )

    except Exception as exc:
        return ExtractedFileResult(
            filename=filename,
            original_size=len(data),
            category="archive",
            mime_hint="application/zip",
            error=f"Archive extraction failed: {exc}",
            summary="Failed to extract archive.",
        )


def process_file_content(
    file_bytes: bytes,
    filename: str,
    sandbox_workspace_dir: Optional[Path] = None,
) -> ExtractedFileResult:
    """Master entrypoint: Validates size (<= 10MB) and processes any file format."""
    size = len(file_bytes)
    if size > MAX_FILE_SIZE_BYTES:
        return ExtractedFileResult(
            filename=filename,
            original_size=size,
            category="binary",
            mime_hint="application/octet-stream",
            error=f"File exceeds maximum allowed limit of 10 MB ({size} bytes).",
            summary="File rejected: size exceeds 10 MB limit.",
        )

    category, ext = get_file_category(filename)

    # 1. Documents: PDF
    if ext == ".pdf":
        return extract_pdf_text(file_bytes, filename)

    # 2. Documents: DOCX
    if ext == ".docx":
        return extract_docx_text(file_bytes, filename)

    # 3. Archives: ZIP, TAR, RAR
    if category == "archive":
        target_dir = (sandbox_workspace_dir or Path(".nexora_uploads")) / f"extracted_{Path(filename).stem}"
        return extract_archive(file_bytes, filename, target_dir)

    # 4. Code, Data, Text Files
    text = decode_text_bytes(file_bytes)
    line_count = len(text.splitlines())
    summary = f"{ext.upper() or 'Text'} file with {line_count} line(s), {size} bytes."

    return ExtractedFileResult(
        filename=filename,
        original_size=size,
        category=category,
        mime_hint=f"text/{ext.lstrip('.') or 'plain'}",
        text_content=text,
        summary=summary,
    )
